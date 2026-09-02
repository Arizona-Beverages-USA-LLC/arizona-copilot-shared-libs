"""
Report document export - render analysis markdown to a formatted PDF or PPTX.

The agent produces analyses as markdown: '##' section headers, '-' insight
bullets, and pipe tables (the required "insight + supporting table" per
section). Charts are placed inline with a marker on its own line:

    [[chart: <chart title>]]     -> match a cached chart by title / id
    [[chart]]                    -> take the next unplaced chart in order

Each marker is replaced by that chart's image AT THAT SPOT in the section.
Charts the model did not place are appended at the end (safety net).

This module renders a professionally formatted PDF (reportlab) or PowerPoint
(python-pptx) and returns bytes for the artifact layer (save_as_adk_artifact).
Pure rendering: no ADK, no network. It takes chart PNG bytes (the tool does
Vega->PNG), so this module has no vl-convert dependency. Brand navy #1F3864.
"""

from __future__ import annotations

import io
import re
from html import escape

_BRAND_HEX = "#1F3864"
_BRAND_RGB = (0x1F, 0x38, 0x64)
_ACCENT_RGB = (0xD9, 0xE1, 0xF2)
_GREY_RGB = (0x59, 0x59, 0x59)
_GRID_HEX = "#BFBFBF"
_ACCENT_HEX = "#D9E1F2"
_DEFAULT_SUBTITLE = "Sales & Merchandising Copilot  |  Arizona Beverages"
_MAX_PPTX_TABLE_ROWS = 16

_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
_CHART_MARKER = re.compile(r"^\[\[\s*chart\s*(?::\s*(.*?))?\s*\]\]$", re.IGNORECASE)


def _split_row(line: str) -> list:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_markdown(md: str) -> list:
    """Parse report markdown into a flat list of blocks: h1/h2/h3, para,
    bullets, table (header, rows), and chart (key)."""
    lines = (md or "").replace("\r\n", "\n").split("\n")
    blocks: list = []
    bullet_buf: list = []
    para_buf: list = []
    n = len(lines)
    i = 0

    def flush_bullets():
        nonlocal bullet_buf
        if bullet_buf:
            blocks.append({"type": "bullets", "items": bullet_buf})
            bullet_buf = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            text = " ".join(x.strip() for x in para_buf).strip()
            if text:
                blocks.append({"type": "para", "text": text})
            para_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        mc = _CHART_MARKER.match(stripped)
        if mc:
            flush_bullets(); flush_para()
            blocks.append({"type": "chart", "key": (mc.group(1) or "").strip()})
            i += 1
            continue

        if "|" in line and i + 1 < n and _TABLE_SEP.match(lines[i + 1]):
            flush_bullets(); flush_para()
            header = _split_row(line)
            rows = []
            i += 2
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if not stripped:
            flush_bullets(); flush_para()
            i += 1
            continue

        mh = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if mh:
            flush_bullets(); flush_para()
            lvl = len(mh.group(1))
            kind = "h1" if lvl == 1 else "h2" if lvl == 2 else "h3"
            blocks.append({"type": kind, "text": mh.group(2).strip()})
            i += 1
            continue

        mb = re.match(r"^[-*+]\s+(.*)$", stripped)
        if mb:
            flush_para()
            bullet_buf.append(mb.group(1).strip())
            i += 1
            continue

        mn = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if mn:
            flush_para()
            bullet_buf.append(mn.group(1).strip())
            i += 1
            continue

        flush_bullets()
        para_buf.append(stripped)
        i += 1

    flush_bullets(); flush_para()
    return blocks


def _strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def _rl_inline(text: str) -> str:
    out = escape(text)
    out = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", out)
    out = re.sub(r"`(.+?)`", r'<font face="Courier">\1</font>', out)
    return out


def _resolve_chart(charts: list, used: set, key: str):
    """Return the chart matching `key` (id, then title-substring, else next
    unplaced) and mark it used. Bare/empty key -> next unplaced in order."""
    k = (key or "").strip().lower()
    if not k:
        for idx, c in enumerate(charts):
            if idx not in used:
                used.add(idx); return c
        return None
    for idx, c in enumerate(charts):
        if idx in used:
            continue
        if str(c.get("id", "")).lower() == k:
            used.add(idx); return c
    for idx, c in enumerate(charts):
        if idx in used:
            continue
        title = str(c.get("title", "")).lower()
        if title and (k in title or title in k):
            used.add(idx); return c
    for idx, c in enumerate(charts):
        if idx not in used:
            used.add(idx); return c
    return None


# ============================================================ PDF (reportlab)
def render_markdown_to_pdf(markdown: str, title: str, subtitle: str = "",
                           charts: list | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, ListFlowable, ListItem,
                                    Image as RLImage)

    brand = colors.HexColor(_BRAND_HEX)
    accent = colors.HexColor(_ACCENT_HEX)
    grey = colors.HexColor("#595959")
    grid = colors.HexColor(_GRID_HEX)
    charts = charts or []
    used: set = set()

    ss = getSampleStyleSheet()
    st = {
        "title": ParagraphStyle("t", parent=ss["Title"], textColor=brand, fontSize=20, spaceAfter=2),
        "subtitle": ParagraphStyle("st", parent=ss["Normal"], textColor=grey, fontSize=10, spaceAfter=14),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], textColor=brand, fontSize=14, spaceBefore=14, spaceAfter=6),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], textColor=brand, fontSize=11, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=13, spaceAfter=6),
        "bullet": ParagraphStyle("bl", parent=ss["Normal"], fontSize=9.5, leading=13),
        "cap": ParagraphStyle("cap", parent=ss["Normal"], fontSize=9, leading=11, textColor=grey, spaceAfter=2),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=8, leading=10),
        "cellh": ParagraphStyle("ch", parent=ss["Normal"], fontSize=8, leading=10,
                                textColor=colors.white, fontName="Helvetica-Bold"),
    }

    story = [Paragraph(escape(title or "Analysis"), st["title"]),
             Paragraph(escape(subtitle or _DEFAULT_SUBTITLE), st["subtitle"])]
    avail = letter[0] - 1.2 * inch

    def emit_chart(ch):
        png = ch.get("png") if isinstance(ch, dict) else None
        if not png:
            return
        story.append(Paragraph(_rl_inline(ch.get("title") or "Chart"), st["cap"]))
        try:
            iw, ih = ImageReader(io.BytesIO(png)).getSize()
            w = avail
            h = w * (ih / iw) if iw else w * 0.5625
            story.append(RLImage(io.BytesIO(png), width=w, height=h))
            story.append(Spacer(1, 12))
        except Exception:
            pass

    for b in parse_markdown(markdown):
        t = b["type"]
        if t in ("h1", "h2"):
            story.append(Paragraph(_rl_inline(b["text"]), st["h2"]))
        elif t == "h3":
            story.append(Paragraph(_rl_inline(b["text"]), st["h3"]))
        elif t == "para":
            story.append(Paragraph(_rl_inline(b["text"]), st["body"]))
        elif t == "bullets":
            items = [ListItem(Paragraph(_rl_inline(x), st["bullet"])) for x in b["items"]]
            story.append(ListFlowable(items, bulletType="bullet", bulletColor=brand,
                                      start="\u2022", leftIndent=14))
            story.append(Spacer(1, 6))
        elif t == "chart":
            ch = _resolve_chart(charts, used, b.get("key"))
            if ch:
                emit_chart(ch)
        elif t == "table":
            header = b["header"]
            ncol = max(1, len(header))
            data = [[Paragraph(_rl_inline(c), st["cellh"]) for c in header]]
            for r in b["rows"]:
                cells = (r + [""] * ncol)[:ncol]
                data.append([Paragraph(_rl_inline(c), st["cell"]) for c in cells])
            tbl = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[avail / ncol] * ncol)
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), brand),
                ("GRID", (0, 0), (-1, -1), 0.4, grid),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
            for ri in range(2, len(data), 2):
                style.append(("BACKGROUND", (0, ri), (-1, ri), accent))
            tbl.setStyle(TableStyle(style))
            story.append(tbl)
            story.append(Spacer(1, 10))

    leftover = [c for idx, c in enumerate(charts) if idx not in used]
    if leftover:
        story.append(Paragraph("Additional Charts", st["h2"]))
        for ch in leftover:
            emit_chart(ch)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title=(title or "Analysis"),
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            leftMargin=0.6 * inch, rightMargin=0.6 * inch)
    doc.build(story)
    return buf.getvalue()


# ============================================================ PPTX (python-pptx)
def _sectionize(blocks: list, deck_title: str) -> list:
    """Group flat blocks into sections. Each section keeps its bullet `lines`
    and an ordered `items` list of ('table', block) / ('chart', key) so tables
    and charts render as slides in document order right after the bullets."""
    sections: list = []
    cur = {"title": deck_title, "lines": [], "items": []}

    def close():
        if cur["lines"] or cur["items"]:
            sections.append(dict(cur))

    for b in blocks:
        t = b["type"]
        if t in ("h1", "h2"):
            close()
            cur = {"title": _strip_md(b["text"]), "lines": [], "items": []}
        elif t == "h3":
            cur["lines"].append(("sub", _strip_md(b["text"])))
        elif t == "para":
            cur["lines"].append(("para", _strip_md(b["text"])))
        elif t == "bullets":
            for it in b["items"]:
                cur["lines"].append(("bullet", _strip_md(it)))
        elif t == "table":
            cur["items"].append(("table", b))
        elif t == "chart":
            cur["items"].append(("chart", b.get("key", "")))
    close()
    return sections


def render_markdown_to_pptx(markdown: str, title: str, subtitle: str = "",
                            charts: list | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    brand = RGBColor(*_BRAND_RGB)
    grey = RGBColor(*_GREY_RGB)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    accent = RGBColor(*_ACCENT_RGB)
    charts = charts or []
    used: set = set()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    SW, SH = prs.slide_width, prs.slide_height

    def title_bar(slide, text):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), SW - Inches(1.0), Inches(0.9))
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = _strip_md(text)
        p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = brand
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.2),
                                      SW - Inches(1.0), Pt(3))
        rule.fill.solid(); rule.fill.fore_color.rgb = brand; rule.line.fill.background()

    def chart_slide(ch):
        png = ch.get("png") if isinstance(ch, dict) else None
        if not png:
            return
        sl = prs.slides.add_slide(blank)
        title_bar(sl, ch.get("title") or "Chart")
        try:
            sl.shapes.add_picture(io.BytesIO(png), Inches(0.7), Inches(1.55), width=SW - Inches(1.4))
        except Exception:
            pass

    # title slide
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.5), SW - Inches(1.6), Inches(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = _strip_md(title or "Analysis")
    p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = brand
    sp = tf.add_paragraph(); sp.text = subtitle or _DEFAULT_SUBTITLE
    sp.font.size = Pt(16); sp.font.color.rgb = grey

    for sec in _sectionize(parse_markdown(markdown), _strip_md(title or "Analysis")):
        if sec["lines"]:
            sl = prs.slides.add_slide(blank); title_bar(sl, sec["title"])
            body = sl.shapes.add_textbox(Inches(0.6), Inches(1.45),
                                         SW - Inches(1.2), SH - Inches(1.95))
            btf = body.text_frame; btf.word_wrap = True
            first = True
            for kind, text in sec["lines"]:
                para = btf.paragraphs[0] if first else btf.add_paragraph()
                first = False
                if kind == "bullet":
                    para.text = "\u2022 " + text; para.font.size = Pt(14)
                elif kind == "sub":
                    para.text = text; para.font.size = Pt(15)
                    para.font.bold = True; para.font.color.rgb = brand
                else:
                    para.text = text; para.font.size = Pt(13)
        for kind, payload in sec["items"]:
            if kind == "table":
                _pptx_table_slide(prs, blank, title_bar, sec["title"], payload,
                                  brand, white, accent, SW, SH)
            else:
                ch = _resolve_chart(charts, used, payload)
                if ch:
                    chart_slide(ch)

    for idx, ch in enumerate(charts):
        if idx not in used:
            chart_slide(ch)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pptx_table_slide(prs, blank, title_bar, section_title, tblblock,
                      brand, white, accent, SW, SH):
    from pptx.util import Inches, Pt

    sl = prs.slides.add_slide(blank)
    header = tblblock["header"] or [""]
    rows = tblblock["rows"]
    truncated = len(rows) > _MAX_PPTX_TABLE_ROWS
    if truncated:
        rows = rows[:_MAX_PPTX_TABLE_ROWS]
    title_bar(sl, section_title + (" (data)" if not truncated
                                   else f" (data - top {_MAX_PPTX_TABLE_ROWS})"))
    ncol = max(1, len(header))
    nrow = len(rows) + 1
    left, top = Inches(0.5), Inches(1.45)
    width = SW - Inches(1.0)
    height = min(SH - Inches(1.8), Inches(0.34) * nrow)
    table = sl.shapes.add_table(nrow, ncol, left, top, width, height).table

    for c, h in enumerate(header):
        cell = table.cell(0, c); cell.text = _strip_md(h)
        pr = cell.text_frame.paragraphs[0]
        pr.font.size = Pt(10); pr.font.bold = True; pr.font.color.rgb = white
        cell.fill.solid(); cell.fill.fore_color.rgb = brand
    for r, row in enumerate(rows, start=1):
        cells = (row + [""] * ncol)[:ncol]
        for c, val in enumerate(cells):
            cell = table.cell(r, c); cell.text = _strip_md(val)
            cell.text_frame.paragraphs[0].font.size = Pt(9)
            if r % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = accent


__all__ = ["parse_markdown", "render_markdown_to_pdf", "render_markdown_to_pptx"]
